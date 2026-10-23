# ============================================================
# ГЛАВНЫЙ ФАЙЛ ВЕБ-ПРИЛОЖЕНИЯ
# ============================================================
from flask import Flask, render_template, request, redirect, url_for, flash, send_file, jsonify, Response
from functools import wraps
from flask_sqlalchemy import SQLAlchemy
from datetime import datetime, date, timedelta
from sqlalchemy import func
from werkzeug.utils import secure_filename
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from io import BytesIO
import os
import re
import pandas as pd

# Геокодер: DaData + Nominatim (резерв)
from geocode import geocode_address, geocode_many

app = Flask(__name__)
app.config['SECRET_KEY'] = 'moya-logistika-secret-key-2026'
# ============================================================
# АВТОРИЗАЦИЯ
# ============================================================
# 
AUTH_USERNAME = "admin"
AUTH_PASSWORD = "Logistika2026!"


@app.before_request
def require_login():
    """
    Защита всех страниц паролем (HTTP Basic Auth).
    Исключения: /static/ (иначе не загрузятся стили и картинки).
    """
    # Статику отдаём без пароля
    if request.path.startswith('/static/'):
        return

    auth = request.authorization
    if not auth or auth.username != AUTH_USERNAME or auth.password != AUTH_PASSWORD:
        return Response(
            '🔒 Требуется авторизация.\n'
            'Введите логин и пароль во всплывающем окне браузера.',
            401,
            {'WWW-Authenticate': 'Basic realm="MyLogistics"'}
        )
# ============================================================

basedir = os.path.abspath(os.path.dirname(__file__))
app.config['SQLALCHEMY_DATABASE_URI'] = 'sqlite:///' + os.path.join(basedir, 'logistics.db')
app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False
app.config['UPLOAD_FOLDER'] = os.path.join(basedir, 'uploads')
app.config['MAX_CONTENT_LENGTH'] = 16 * 1024 * 1024

if not os.path.exists(app.config['UPLOAD_FOLDER']):
    os.makedirs(app.config['UPLOAD_FOLDER'])

db = SQLAlchemy(app)

KEG_PRICE = 2000

# ============================================================
# МОДЕЛИ
# ============================================================

class Client(db.Model):
    __tablename__ = 'clients'
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(200), nullable=False)
    address = db.Column(db.String(500))
    phone = db.Column(db.String(50))
    contact_person = db.Column(db.String(200))
    group_abc = db.Column(db.String(1), default='C')
    comment = db.Column(db.Text)
    created_at = db.Column(db.DateTime, default=datetime.now)
    lat = db.Column(db.Float)
    lon = db.Column(db.Float)
    orders = db.relationship('Order', backref='client', lazy=True)
    kegs = db.relationship('Keg', backref='client', lazy=True)


class Driver(db.Model):
    __tablename__ = 'drivers'
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(200), nullable=False)
    phone = db.Column(db.String(50))
    vehicle = db.Column(db.String(100))
    status = db.Column(db.String(30), default='Активен')
    comment = db.Column(db.Text)
    created_at = db.Column(db.DateTime, default=datetime.now)
    routes = db.relationship('Route', backref='driver', lazy=True)


class Route(db.Model):
    __tablename__ = 'routes'
    id = db.Column(db.Integer, primary_key=True)
    driver_id = db.Column(db.Integer, db.ForeignKey('drivers.id'), nullable=False)
    route_date = db.Column(db.Date, default=date.today)
    status = db.Column(db.String(30), default='Новый')
    comment = db.Column(db.Text)
    created_at = db.Column(db.DateTime, default=datetime.now)
    orders = db.relationship('Order', backref='route', lazy=True)


class Order(db.Model):
    __tablename__ = 'orders'
    id = db.Column(db.Integer, primary_key=True)
    client_id = db.Column(db.Integer, db.ForeignKey('clients.id'), nullable=False)
    route_id = db.Column(db.Integer, db.ForeignKey('routes.id'), nullable=True)
    order_date = db.Column(db.Date, default=date.today)
    delivery_date = db.Column(db.Date)
    amount = db.Column(db.Float, default=0)
    status = db.Column(db.String(30), default='Новый')
    comment = db.Column(db.Text)
    created_at = db.Column(db.DateTime, default=datetime.now)


class Keg(db.Model):
    __tablename__ = 'kegs'
    id = db.Column(db.Integer, primary_key=True)
    client_id = db.Column(db.Integer, db.ForeignKey('clients.id'), nullable=False)
    operation_date = db.Column(db.Date, default=date.today)
    operation_type = db.Column(db.String(20), default='Выдано')
    quantity = db.Column(db.Integer, default=0)
    comment = db.Column(db.Text)
    created_at = db.Column(db.DateTime, default=datetime.now)


class ImportLog(db.Model):
    __tablename__ = 'import_logs'
    id = db.Column(db.Integer, primary_key=True)
    filename = db.Column(db.String(500))
    file_hash = db.Column(db.String(64))
    order_date = db.Column(db.Date)
    clients_created = db.Column(db.Integer, default=0)
    clients_found = db.Column(db.Integer, default=0)
    orders_created = db.Column(db.Integer, default=0)
    orders_skipped = db.Column(db.Integer, default=0)
    total_amount = db.Column(db.Float, default=0)
    created_at = db.Column(db.DateTime, default=datetime.now)


with app.app_context():
    db.create_all()
    print("✅ База данных готова")

# ============================================================
# ВСПОМОГАТЕЛЬНЫЕ ФУНКЦИИ
# ============================================================
def normalize_name(name):
    if not name:
        return ''
    name = str(name).lower().strip()
    name = re.sub(r'\b(ип|ооо|оао|зао|пао|индивидуальный предприниматель)\b', '', name)
    name = re.sub(r'[.,"«»\'()\[\]]', '', name)
    name = re.sub(r'\s+', ' ', name).strip()
    return name


def find_client_by_name(name):
    if not name:
        return None
    normalized = normalize_name(name)
    for client in Client.query.all():
        if normalize_name(client.name) == normalized:
            return client
    return None


def order_exists(client_id, amount, order_date):
    return Order.query.filter_by(
        client_id=client_id, amount=amount, order_date=order_date
    ).first() is not None


def get_file_hash(filepath):
    import hashlib
    md5 = hashlib.md5()
    with open(filepath, 'rb') as f:
        for chunk in iter(lambda: f.read(4096), b''):
            md5.update(chunk)
    return md5.hexdigest()


def try_autogeocode(client):
    """
    Пытается заполнить lat/lon клиента через DaData, если их нет.
    Возвращает True, если координаты были установлены.
    """
    if client.lat is not None and client.lon is not None:
        return False
    if not client.address:
        return False
    try:
        res = geocode_address(client.address, city=None)
    except Exception as e:
        print(f'[autogeocode] ошибка для клиента {client.id}: {e}')
        return False
    if not res:
        return False
    client.lat = res['lat']
    client.lon = res['lon']
    return True


def style_header(ws, row, columns):
    header_fill = PatternFill(start_color='667EEA', end_color='667EEA', fill_type='solid')
    header_font = Font(bold=True, color='FFFFFF', size=11)
    border = Border(left=Side(style='thin', color='CCCCCC'),
                    right=Side(style='thin', color='CCCCCC'),
                    top=Side(style='thin', color='CCCCCC'),
                    bottom=Side(style='thin', color='CCCCCC'))
    for col_idx, col_name in enumerate(columns, 1):
        cell = ws.cell(row=row, column=col_idx, value=col_name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = border


def style_cell(cell, align='left'):
    border = Border(left=Side(style='thin', color='EEEEEE'),
                    right=Side(style='thin', color='EEEEEE'),
                    top=Side(style='thin', color='EEEEEE'),
                    bottom=Side(style='thin', color='EEEEEE'))
    cell.border = border
    cell.alignment = Alignment(horizontal=align, vertical='center', wrap_text=True)


def auto_width(ws, min_width=10, max_width=50):
    for col in ws.columns:
        max_length = 0
        col_letter = None
        for cell in col:
            if not hasattr(cell, 'column_letter'):
                continue
            if col_letter is None:
                col_letter = cell.column_letter
            try:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            except:
                pass
        if col_letter:
            adjusted = min(max(max_length + 2, min_width), max_width)
            ws.column_dimensions[col_letter].width = adjusted


# ============================================================
# ГЛАВНАЯ
# ============================================================
@app.route('/')
def index():
    keg_debt = db.session.query(
        func.sum(Keg.quantity).filter(Keg.operation_type == 'Выдано')
    ).scalar() or 0
    keg_returned = db.session.query(
        func.sum(Keg.quantity).filter(Keg.operation_type == 'Возвращено')
    ).scalar() or 0
    keg_debt_total = keg_debt - keg_returned

    clients_no_coords = Client.query.filter(
        (Client.lat.is_(None)) | (Client.lon.is_(None))
    ).count()

    return render_template('index.html',
                           clients_count=Client.query.count(),
                           orders_count=Order.query.count(),
                           drivers_count=Driver.query.count(),
                           routes_count=Route.query.count(),
                           keg_debt_total=keg_debt_total,
                           clients_no_coords=clients_no_coords)

# ============================================================
# КЛИЕНТЫ
# ============================================================
@app.route('/clients')
def clients():
    return render_template('clients.html', clients=Client.query.order_by(Client.name).all())

@app.route('/clients/add', methods=['GET', 'POST'])
def client_add():
    if request.method == 'POST':
        name = request.form.get('name', '').strip()
        if not name:
            flash('Имя клиента обязательно!', 'error')
            return redirect(url_for('client_add'))

        lat_str = request.form.get('lat', '').strip()
        lon_str = request.form.get('lon', '').strip()
        try:
            lat_val = float(lat_str) if lat_str else None
        except ValueError:
            lat_val = None
        try:
            lon_val = float(lon_str) if lon_str else None
        except ValueError:
            lon_val = None

        client = Client(
            name=name, address=request.form.get('address', '').strip(),
            phone=request.form.get('phone', '').strip(),
            contact_person=request.form.get('contact_person', '').strip(),
            group_abc=request.form.get('group_abc', 'C').strip(),
            comment=request.form.get('comment', '').strip(),
            lat=lat_val, lon=lon_val)

        # Авто-геокодирование, если координаты не заданы вручную
        db.session.add(client)
        db.session.flush()  # получаем client.id
        if try_autogeocode(client):
            flash(f'📍 Координаты для "{name}" подтянуты автоматически.', 'success')

        db.session.commit()
        flash(f'Клиент "{name}" добавлен!', 'success')
        return redirect(url_for('clients'))
    return render_template('client_add.html')

@app.route('/clients/<int:client_id>')
def client_view(client_id):
    client = Client.query.get_or_404(client_id)
    keg_debt = db.session.query(func.sum(Keg.quantity)).filter(
        Keg.client_id == client_id, Keg.operation_type == 'Выдано'
    ).scalar() or 0
    keg_returned = db.session.query(func.sum(Keg.quantity)).filter(
        Keg.client_id == client_id, Keg.operation_type == 'Возвращено'
    ).scalar() or 0
    keg_balance = keg_debt - keg_returned
    return render_template('client_view.html', client=client, keg_balance=keg_balance)

@app.route('/clients/<int:client_id>/edit', methods=['GET', 'POST'])
def client_edit(client_id):
    client = Client.query.get_or_404(client_id)
    if request.method == 'POST':
        client.name = request.form.get('name', '').strip()
        client.address = request.form.get('address', '').strip()
        client.phone = request.form.get('phone', '').strip()
        client.contact_person = request.form.get('contact_person', '').strip()
        client.group_abc = request.form.get('group_abc', 'C').strip()
        client.comment = request.form.get('comment', '').strip()

        lat_str = request.form.get('lat', '').strip()
        lon_str = request.form.get('lon', '').strip()
        try:
            client.lat = float(lat_str) if lat_str else None
        except ValueError:
            client.lat = None
        try:
            client.lon = float(lon_str) if lon_str else None
        except ValueError:
            client.lon = None

        if not client.name:
            flash('Имя обязательно!', 'error')
            return redirect(url_for('client_edit', client_id=client_id))

        # Авто-геокодирование, если адрес есть, а координат нет
        if try_autogeocode(client):
            flash('📍 Координаты подтянуты автоматически по адресу.', 'success')

        db.session.commit()
        flash(f'Клиент "{client.name}" обновлён!', 'success')
        return redirect(url_for('client_view', client_id=client_id))
    return render_template('client_edit.html', client=client)

@app.route('/clients/<int:client_id>/delete', methods=['POST'])
def client_delete(client_id):
    client = Client.query.get_or_404(client_id)
    name = client.name
    db.session.delete(client)
    db.session.commit()
    flash(f'Клиент "{name}" удалён!', 'success')
    return redirect(url_for('clients'))

@app.route('/clients/no_coords')
def clients_no_coords():
    """Список клиентов без координат"""
    clients_list = Client.query.filter(
        (Client.lat.is_(None)) | (Client.lon.is_(None))
    ).order_by(Client.name).all()
    return render_template('clients_no_coords.html', clients=clients_list)

# ------------------------------------------------------------
# ГЕОКОДИРОВАНИЕ: API и массовая операция
# ------------------------------------------------------------
@app.route('/clients/no_coords/geocode_all', methods=['POST'])
def clients_no_coords_geocode_all():
    """Прогоняет DaData по всем клиентам без координат."""
    clients_list = Client.query.filter(
        (Client.lat.is_(None)) | (Client.lon.is_(None))
    ).all()
    items = [(c.id, c.address) for c in clients_list if c.address]
    if not items:
        flash('Нет клиентов с адресом для геокодирования.', 'error')
        return redirect(url_for('clients_no_coords'))

    try:
        results = geocode_many(items, city=None, pause=0.2)
    except Exception as e:
        print(f'[geocode_all] общая ошибка: {e}')
        flash(f'Ошибка геокодирования: {e}', 'error')
        return redirect(url_for('clients_no_coords'))

    done = 0
    for r in results:
        c = Client.query.get(r['id'])
        if c:
            c.lat = r['lat']
            c.lon = r['lon']
            done += 1
    db.session.commit()

    failed = len(items) - done
    msg = f'🌍 Геокодирование завершено: успешно {done}'
    if failed > 0:
        msg += f', не найдено {failed}'
    flash(msg, 'success')
    return redirect(url_for('clients_no_coords'))


@app.route('/api/geocode', methods=['POST'])
def api_geocode():
    """Принимает {"address": "..."} → возвращает {ok, lat, lon, address, source}."""
    data = request.get_json(silent=True) or {}
    address = (data.get('address') or '').strip()
    if not address:
        return jsonify({'ok': False, 'error': 'Пустой адрес'}), 400
    try:
        res = geocode_address(address, city=None)
    except Exception as e:
        return jsonify({'ok': False, 'error': str(e)}), 500
    if not res:
        return jsonify({'ok': False, 'error': 'Адрес не найден'})
    return jsonify({
        'ok': True,
        'lat': res['lat'],
        'lon': res['lon'],
        'address': res.get('address', address),
        'source': res.get('source', 'unknown'),
    })


@app.route('/api/clients/<int:client_id>/coords', methods=['POST'])
def api_client_coords(client_id):
    """Принимает {"lat": ..., "lon": ...}, сохраняет в БД."""
    client = Client.query.get_or_404(client_id)
    data = request.get_json(silent=True) or {}
    try:
        lat = float(data.get('lat'))
        lon = float(data.get('lon'))
    except (TypeError, ValueError):
        return jsonify({'ok': False, 'error': 'Некорректные координаты'}), 400
    client.lat = lat
    client.lon = lon
    db.session.commit()
    return jsonify({'ok': True, 'lat': lat, 'lon': lon})


@app.route('/export/clients')
def export_clients():
    wb = Workbook()
    ws = wb.active
    ws.title = 'Клиенты'
    columns = ['ID', 'Название', 'Адрес', 'Телефон', 'Контактное лицо', 'Группа ABC', 'Широта', 'Долгота', 'Комментарий', 'Добавлен']
    style_header(ws, 1, columns)
    clients_list = Client.query.order_by(Client.name).all()
    for row_idx, client in enumerate(clients_list, 2):
        ws.cell(row=row_idx, column=1, value=client.id)
        ws.cell(row=row_idx, column=2, value=client.name)
        ws.cell(row=row_idx, column=3, value=client.address or '')
        ws.cell(row=row_idx, column=4, value=client.phone or '')
        ws.cell(row=row_idx, column=5, value=client.contact_person or '')
        ws.cell(row=row_idx, column=6, value=client.group_abc or 'C')
        ws.cell(row=row_idx, column=7, value=client.lat or '')
        ws.cell(row=row_idx, column=8, value=client.lon or '')
        ws.cell(row=row_idx, column=9, value=client.comment or '')
        ws.cell(row=row_idx, column=10, value=client.created_at.strftime('%d.%m.%Y %H:%M') if client.created_at else '')
        for col_idx in range(1, 11):
            style_cell(ws.cell(row=row_idx, column=col_idx), 'center' if col_idx in [1, 6, 7, 8, 10] else 'left')
    auto_width(ws)
    ws.freeze_panes = 'A2'
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    filename = f'Клиенты_{datetime.now().strftime("%d.%m.%Y_%H-%M")}.xlsx'
    return send_file(output, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     as_attachment=True, download_name=filename)

# ============================================================
# ЗАКАЗЫ
# ============================================================
@app.route('/orders')
def orders():
    status_filter = request.args.get('status', '').strip()
    client_filter = request.args.get('client_id', '').strip()
    date_from = request.args.get('date_from', '').strip()
    date_to = request.args.get('date_to', '').strip()
    search = request.args.get('search', '').strip()
    route_filter = request.args.get('route_id', '').strip()
    page = request.args.get('page', 1, type=int)
    per_page = request.args.get('per_page', 50, type=int)

    query = Order.query
    if status_filter:
        query = query.filter(Order.status == status_filter)
    if client_filter:
        try:
            query = query.filter(Order.client_id == int(client_filter))
        except ValueError:
            pass
    if date_from:
        try:
            d = datetime.strptime(date_from, '%Y-%m-%d').date()
            query = query.filter(Order.order_date >= d)
        except ValueError:
            pass
    if date_to:
        try:
            d = datetime.strptime(date_to, '%Y-%m-%d').date()
            query = query.filter(Order.order_date <= d)
        except ValueError:
            pass
    if route_filter == 'none':
        query = query.filter(Order.route_id.is_(None))
    elif route_filter:
        try:
            query = query.filter(Order.route_id == int(route_filter))
        except ValueError:
            pass
    if search:
        try:
            search_id = int(search)
            query = query.filter(Order.id == search_id)
        except ValueError:
            query = query.join(Client).filter(Client.name.ilike(f'%{search}%'))

    query = query.order_by(Order.order_date.desc(), Order.id.desc())
    pagination = query.paginate(page=page, per_page=per_page, error_out=False)
    orders_list = pagination.items

    stats = {
        'Новый': Order.query.filter_by(status='Новый').count(),
        'В работе': Order.query.filter_by(status='В работе').count(),
        'Доставлен': Order.query.filter_by(status='Доставлен').count(),
        'Отменён': Order.query.filter_by(status='Отменён').count(),
        'всего': Order.query.count(),
        'без маршрута': Order.query.filter(Order.route_id.is_(None)).count(),
        'показано': pagination.total,
    }

    all_clients = Client.query.order_by(Client.name).all()
    all_routes = Route.query.order_by(Route.route_date.desc()).limit(50).all()

    return render_template('orders.html', orders=orders_list, pagination=pagination,
                           stats=stats, clients=all_clients, routes=all_routes,
                           current_filter=status_filter, current_client=client_filter,
                           current_date_from=date_from, current_date_to=date_to,
                           current_search=search, current_route=route_filter, per_page=per_page)


@app.route('/export/orders')
def export_orders():
    status_filter = request.args.get('status', '').strip()
    client_filter = request.args.get('client_id', '').strip()
    date_from = request.args.get('date_from', '').strip()
    date_to = request.args.get('date_to', '').strip()
    search = request.args.get('search', '').strip()
    route_filter = request.args.get('route_id', '').strip()

    query = Order.query
    if status_filter:
        query = query.filter(Order.status == status_filter)
    if client_filter:
        try:
            query = query.filter(Order.client_id == int(client_filter))
        except ValueError:
            pass
    if date_from:
        try:
            query = query.filter(Order.order_date >= datetime.strptime(date_from, '%Y-%m-%d').date())
        except ValueError:
            pass
    if date_to:
        try:
            query = query.filter(Order.order_date <= datetime.strptime(date_to, '%Y-%m-%d').date())
        except ValueError:
            pass
    if route_filter == 'none':
        query = query.filter(Order.route_id.is_(None))
    elif route_filter:
        try:
            query = query.filter(Order.route_id == int(route_filter))
        except ValueError:
            pass
    if search:
        try:
            query = query.filter(Order.id == int(search))
        except ValueError:
            query = query.join(Client).filter(Client.name.ilike(f'%{search}%'))

    query = query.order_by(Order.order_date.desc(), Order.id.desc())
    orders_list = query.all()

    wb = Workbook()
    ws = wb.active
    ws.title = 'Заказы'
    columns = ['№', 'Дата', 'Клиент', 'Адрес', 'Телефон', 'Сумма (₽)', 'Статус', 'Маршрут', 'Комментарий']
    style_header(ws, 1, columns)
    total_amount = 0
    for row_idx, order in enumerate(orders_list, 2):
        client = order.client
        ws.cell(row=row_idx, column=1, value=order.id)
        ws.cell(row=row_idx, column=2, value=order.order_date.strftime('%d.%m.%Y') if order.order_date else '')
        ws.cell(row=row_idx, column=3, value=client.name if client else '')
        ws.cell(row=row_idx, column=4, value=client.address if client and client.address else '')
        ws.cell(row=row_idx, column=5, value=client.phone if client and client.phone else '')
        ws.cell(row=row_idx, column=6, value=order.amount)
        ws.cell(row=row_idx, column=7, value=order.status)
        ws.cell(row=row_idx, column=8, value=f'#{order.route_id}' if order.route_id else '—')
        ws.cell(row=row_idx, column=9, value=order.comment or '')
        total_amount += order.amount or 0
        for col_idx in range(1, 10):
            style_cell(ws.cell(row=row_idx, column=col_idx), 'center' if col_idx in [1, 2, 6, 7, 8] else 'left')
    last_row = len(orders_list) + 2
    ws.cell(row=last_row, column=1, value='ИТОГО:').font = Font(bold=True)
    ws.cell(row=last_row, column=6, value=total_amount).font = Font(bold=True)
    ws.cell(row=last_row, column=6).fill = PatternFill(start_color='E8F5E9', end_color='E8F5E9', fill_type='solid')
    auto_width(ws)
    ws.freeze_panes = 'A2'
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    filename = f'Заказы_{datetime.now().strftime("%d.%m.%Y_%H-%M")}.xlsx'
    return send_file(output, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     as_attachment=True, download_name=filename)


@app.route('/orders/mass_action', methods=['POST'])
def orders_mass_action():
    action = request.form.get('action', '')
    order_ids = request.form.getlist('order_ids')
    if not order_ids:
        flash('Не выбрано ни одного заказа!', 'error')
        return redirect(url_for('orders'))
    try:
        ids = [int(oid) for oid in order_ids]
    except ValueError:
        flash('Некорректные ID заказов.', 'error')
        return redirect(url_for('orders'))
    selected_orders = Order.query.filter(Order.id.in_(ids)).all()
    if not selected_orders:
        flash('Заказы не найдены.', 'error')
        return redirect(url_for('orders'))
    count = len(selected_orders)

    if action == 'change_status':
        new_status = request.form.get('new_status', '').strip()
        if not new_status:
            flash('Выберите новый статус!', 'error')
            return redirect(url_for('orders'))
        for order in selected_orders:
            order.status = new_status
        db.session.commit()
        flash(f'✅ Статус изменён на «{new_status}» для {count} заказов.', 'success')
    elif action == 'add_to_route':
        route_id = request.form.get('route_id', '').strip()
        if not route_id:
            flash('Выберите маршрут!', 'error')
            return redirect(url_for('orders'))
        try:
            route = Route.query.get_or_404(int(route_id))
        except ValueError:
            flash('Некорректный маршрут.', 'error')
            return redirect(url_for('orders'))
        for order in selected_orders:
            order.route_id = route.id
            if order.status == 'Новый':
                order.status = 'В работе'
        db.session.commit()
        flash(f'✅ {count} заказов добавлено в маршрут №{route.id}.', 'success')
    elif action == 'remove_from_route':
        for order in selected_orders:
            order.route_id = None
        db.session.commit()
        flash(f'✅ {count} заказов убрано из маршрутов.', 'success')
    elif action == 'delete':
        for order in selected_orders:
            db.session.delete(order)
        db.session.commit()
        flash(f'🗑 {count} заказов удалено.', 'success')
    else:
        flash('Неизвестное действие.', 'error')
    return redirect(request.referrer or url_for('orders'))


@app.route('/orders/add', methods=['GET', 'POST'])
def order_add():
    if request.method == 'POST':
        client_id = request.form.get('client_id', '').strip()
        if not client_id:
            flash('Выберите клиента!', 'error')
            return redirect(url_for('order_add'))
        order_date = request.form.get('order_date', '').strip()
        delivery_date = request.form.get('delivery_date', '').strip()
        try:
            order_date_val = datetime.strptime(order_date, '%Y-%m-%d').date() if order_date else date.today()
        except ValueError:
            order_date_val = date.today()
        try:
            delivery_date_val = datetime.strptime(delivery_date, '%Y-%m-%d').date() if delivery_date else None
        except ValueError:
            delivery_date_val = None
        try:
            amount_val = float(request.form.get('amount', 0))
        except ValueError:
            amount_val = 0
        order = Order(client_id=int(client_id), order_date=order_date_val,
                      delivery_date=delivery_date_val, amount=amount_val,
                      status=request.form.get('status', 'Новый').strip(),
                      comment=request.form.get('comment', '').strip())
        db.session.add(order)
        db.session.commit()
        flash(f'Заказ №{order.id} добавлен!', 'success')
        return redirect(url_for('orders'))
    return render_template('order_add.html', clients=Client.query.order_by(Client.name).all(),
                           today=date.today().strftime('%Y-%m-%d'))

@app.route('/orders/<int:order_id>')
def order_view(order_id):
    return render_template('order_view.html', order=Order.query.get_or_404(order_id))

@app.route('/orders/<int:order_id>/edit', methods=['GET', 'POST'])
def order_edit(order_id):
    order = Order.query.get_or_404(order_id)
    if request.method == 'POST':
        order.client_id = int(request.form.get('client_id', order.client_id))
        order_date = request.form.get('order_date', '')
        try:
            order.order_date = datetime.strptime(order_date, '%Y-%m-%d').date() if order_date else order.order_date
        except ValueError:
            pass
        delivery_date = request.form.get('delivery_date', '')
        try:
            order.delivery_date = datetime.strptime(delivery_date, '%Y-%m-%d').date() if delivery_date else None
        except ValueError:
            order.delivery_date = None
        try:
            order.amount = float(request.form.get('amount', 0))
        except ValueError:
            order.amount = 0
        order.status = request.form.get('status', order.status)
        order.comment = request.form.get('comment', '').strip()
        db.session.commit()
        flash(f'Заказ №{order.id} обновлён!', 'success')
        return redirect(url_for('order_view', order_id=order.id))
    return render_template('order_edit.html', order=order, clients=Client.query.order_by(Client.name).all())

@app.route('/orders/<int:order_id>/delete', methods=['POST'])
def order_delete(order_id):
    order = Order.query.get_or_404(order_id)
    num = order.id
    db.session.delete(order)
    db.session.commit()
    flash(f'Заказ №{num} удалён!', 'success')
    return redirect(url_for('orders'))

# ============================================================
# ВОДИТЕЛИ
# ============================================================
@app.route('/drivers')
def drivers():
    return render_template('drivers.html', drivers=Driver.query.order_by(Driver.name).all())

@app.route('/drivers/add', methods=['GET', 'POST'])
def driver_add():
    if request.method == 'POST':
        name = request.form.get('name', '').strip()
        if not name:
            flash('Имя обязательно!', 'error')
            return redirect(url_for('driver_add'))
        driver = Driver(name=name, phone=request.form.get('phone', '').strip(),
                        vehicle=request.form.get('vehicle', '').strip(),
                        status=request.form.get('status', 'Активен').strip(),
                        comment=request.form.get('comment', '').strip())
        db.session.add(driver)
        db.session.commit()
        flash(f'Водитель "{name}" добавлен!', 'success')
        return redirect(url_for('drivers'))
    return render_template('driver_add.html')

@app.route('/drivers/<int:driver_id>')
def driver_view(driver_id):
    return render_template('driver_view.html', driver=Driver.query.get_or_404(driver_id))

@app.route('/drivers/<int:driver_id>/edit', methods=['GET', 'POST'])
def driver_edit(driver_id):
    driver = Driver.query.get_or_404(driver_id)
    if request.method == 'POST':
        driver.name = request.form.get('name', '').strip()
        driver.phone = request.form.get('phone', '').strip()
        driver.vehicle = request.form.get('vehicle', '').strip()
        driver.status = request.form.get('status', 'Активен').strip()
        driver.comment = request.form.get('comment', '').strip()
        if not driver.name:
            flash('Имя обязательно!', 'error')
            return redirect(url_for('driver_edit', driver_id=driver_id))
        db.session.commit()
        flash(f'Водитель "{driver.name}" обновлён!', 'success')
        return redirect(url_for('driver_view', driver_id=driver_id))
    return render_template('driver_edit.html', driver=driver)

@app.route('/drivers/<int:driver_id>/delete', methods=['POST'])
def driver_delete(driver_id):
    driver = Driver.query.get_or_404(driver_id)
    name = driver.name
    db.session.delete(driver)
    db.session.commit()
    flash(f'Водитель "{name}" удалён!', 'success')
    return redirect(url_for('drivers'))

# ============================================================
# МАРШРУТЫ
# ============================================================
@app.route('/routes')
def routes():
    status_filter = request.args.get('status', '')
    query = Route.query
    if status_filter:
        query = query.filter(Route.status == status_filter)
    all_routes = query.order_by(Route.route_date.desc()).all()
    stats = {
        'Новый': Route.query.filter_by(status='Новый').count(),
        'В работе': Route.query.filter_by(status='В работе').count(),
        'Выполнен': Route.query.filter_by(status='Выполнен').count(),
        'всего': Route.query.count(),
    }
    return render_template('routes.html', routes=all_routes, stats=stats, current_filter=status_filter)

@app.route('/routes/add', methods=['GET', 'POST'])
def route_add():
    if request.method == 'POST':
        driver_id = request.form.get('driver_id', '').strip()
        if not driver_id:
            flash('Выберите водителя!', 'error')
            return redirect(url_for('route_add'))
        route_date = request.form.get('route_date', '').strip()
        try:
            route_date_val = datetime.strptime(route_date, '%Y-%m-%d').date() if route_date else date.today()
        except ValueError:
            route_date_val = date.today()
        route = Route(driver_id=int(driver_id), route_date=route_date_val,
                      status=request.form.get('status', 'Новый').strip(),
                      comment=request.form.get('comment', '').strip())
        db.session.add(route)
        db.session.commit()
        flash(f'Маршрут №{route.id} создан!', 'success')
        return redirect(url_for('route_view', route_id=route.id))
    return render_template('route_add.html',
                           drivers=Driver.query.filter_by(status='Активен').order_by(Driver.name).all(),
                           today=date.today().strftime('%Y-%m-%d'))

@app.route('/routes/<int:route_id>')
def route_view(route_id):
    route = Route.query.get_or_404(route_id)
    free_orders = Order.query.filter(Order.route_id.is_(None)).order_by(Order.order_date).all()
    return render_template('route_view.html', route=route, free_orders=free_orders)

@app.route('/routes/<int:route_id>/print')
def route_print(route_id):
    route = Route.query.get_or_404(route_id)
    return render_template('route_print.html', route=route)

@app.route('/export/route/<int:route_id>')
def export_route(route_id):
    route = Route.query.get_or_404(route_id)
    wb = Workbook()
    ws = wb.active
    ws.title = f'Маршрут {route.id}'
    ws.merge_cells('A1:H1')
    title_cell = ws.cell(row=1, column=1, value=f'МАРШРУТНЫЙ ЛИСТ №{route.id}')
    title_cell.font = Font(bold=True, size=16)
    title_cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 25
    ws.cell(row=3, column=1, value='Водитель:').font = Font(bold=True)
    ws.cell(row=3, column=2, value=route.driver.name if route.driver else '—')
    ws.cell(row=4, column=1, value='Машина:').font = Font(bold=True)
    ws.cell(row=4, column=2, value=route.driver.vehicle if route.driver and route.driver.vehicle else '—')
    ws.cell(row=5, column=1, value='Телефон:').font = Font(bold=True)
    ws.cell(row=5, column=2, value=route.driver.phone if route.driver and route.driver.phone else '—')
    ws.cell(row=3, column=4, value='Дата:').font = Font(bold=True)
    ws.cell(row=3, column=5, value=route.route_date.strftime('%d.%m.%Y') if route.route_date else '—')
    ws.cell(row=4, column=4, value='Статус:').font = Font(bold=True)
    ws.cell(row=4, column=5, value=route.status)
    ws.cell(row=5, column=4, value='Точек:').font = Font(bold=True)
    ws.cell(row=5, column=5, value=len(route.orders))
    ws.cell(row=5, column=6, value='Сумма:').font = Font(bold=True)
    ws.cell(row=5, column=7, value=sum(o.amount for o in route.orders))
    columns = ['№', 'Клиент', 'Адрес', 'Телефон', 'Сумма (₽)', 'Статус', 'Отметка', 'Подпись']
    style_header(ws, 7, columns)
    for row_idx, order in enumerate(route.orders, 8):
        client = order.client
        ws.cell(row=row_idx, column=1, value=row_idx - 7)
        ws.cell(row=row_idx, column=2, value=client.name if client else '')
        ws.cell(row=row_idx, column=3, value=client.address if client and client.address else '')
        ws.cell(row=row_idx, column=4, value=client.phone if client and client.phone else '')
        ws.cell(row=row_idx, column=5, value=order.amount)
        ws.cell(row=row_idx, column=6, value=order.status)
        ws.cell(row=row_idx, column=7, value='')
        ws.cell(row=row_idx, column=8, value='')
        for col_idx in range(1, 9):
            style_cell(ws.cell(row=row_idx, column=col_idx), 'center' if col_idx in [1, 5, 6] else 'left')
        ws.row_dimensions[row_idx].height = 25
    auto_width(ws)
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    filename = f'Маршрут_{route.id}_{datetime.now().strftime("%d.%m.%Y")}.xlsx'
    return send_file(output, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     as_attachment=True, download_name=filename)


@app.route('/routes/<int:route_id>/add_order/<int:order_id>', methods=['POST'])
def route_add_order(route_id, order_id):
    route = Route.query.get_or_404(route_id)
    order = Order.query.get_or_404(order_id)
    order.route_id = route.id
    db.session.commit()
    flash(f'Заказ №{order.id} добавлен в маршрут №{route.id}!', 'success')
    return redirect(url_for('route_view', route_id=route.id))

@app.route('/routes/<int:route_id>/remove_order/<int:order_id>', methods=['POST'])
def route_remove_order(route_id, order_id):
    order = Order.query.get_or_404(order_id)
    order.route_id = None
    db.session.commit()
    flash(f'Заказ №{order.id} убран из маршрута.', 'success')
    return redirect(url_for('route_view', route_id=route_id))

@app.route('/routes/<int:route_id>/edit', methods=['GET', 'POST'])
def route_edit(route_id):
    route = Route.query.get_or_404(route_id)
    if request.method == 'POST':
        route.driver_id = int(request.form.get('driver_id', route.driver_id))
        route_date = request.form.get('route_date', '')
        try:
            route.route_date = datetime.strptime(route_date, '%Y-%m-%d').date() if route_date else route.route_date
        except ValueError:
            pass
        route.status = request.form.get('status', route.status)
        route.comment = request.form.get('comment', '').strip()
        db.session.commit()
        flash(f'Маршрут №{route.id} обновлён!', 'success')
        return redirect(url_for('route_view', route_id=route.id))
    return render_template('route_edit.html', route=route,
                           drivers=Driver.query.filter_by(status='Активен').order_by(Driver.name).all())

@app.route('/routes/<int:route_id>/delete', methods=['POST'])
def route_delete(route_id):
    route = Route.query.get_or_404(route_id)
    for order in route.orders:
        order.route_id = None
    db.session.delete(route)
    db.session.commit()
    flash(f'Маршрут №{route_id} удалён!', 'success')
    return redirect(url_for('routes'))

# ============================================================
# МОБИЛЬНАЯ ВЕРСИЯ
# ============================================================
@app.route('/driver')
def driver_entry():
    active_drivers = Driver.query.filter_by(status='Активен').order_by(Driver.name).all()
    return render_template('driver_entry.html', drivers=active_drivers)


@app.route('/driver/<int:driver_id>')
def driver_home(driver_id):
    driver = Driver.query.get_or_404(driver_id)
    active_routes = Route.query.filter(Route.driver_id == driver_id, Route.status != 'Выполнен')\
        .order_by(Route.route_date.desc()).all()
    done_routes = Route.query.filter(Route.driver_id == driver_id, Route.status == 'Выполнен')\
        .order_by(Route.route_date.desc()).limit(5).all()
    return render_template('driver_home.html', driver=driver,
                           active_routes=active_routes, done_routes=done_routes)


@app.route('/driver/route/<int:route_id>')
def driver_route(route_id):
    route = Route.query.get_or_404(route_id)
    total = len(route.orders)
    done = sum(1 for o in route.orders if o.status == 'Доставлен')
    waiting = total - done

    points = []
    no_coords = []
    for idx, order in enumerate(route.orders, 1):
        client = order.client
        if client and client.lat and client.lon:
            points.append({
                'num': len(points) + 1,
                'order_id': order.id,
                'name': client.name,
                'address': client.address or '',
                'phone': client.phone or '',
                'lat': client.lat,
                'lon': client.lon,
                'amount': order.amount,
                'status': order.status,
            })
        else:
            no_coords.append({
                'order_id': order.id,
                'client_id': client.id if client else None,
                'name': client.name if client else '—',
                'address': client.address if client else '—',
            })

    return render_template('driver_route.html',
                           route=route,
                           total=total,
                           done=done,
                           waiting=waiting,
                           points=points,
                           no_coords=no_coords)


@app.route('/driver/order/<int:order_id>/done', methods=['POST'])
def driver_order_done(order_id):
    order = Order.query.get_or_404(order_id)
    route_id = order.route_id
    order.status = 'Доставлен'
    if route_id:
        route = Route.query.get(route_id)
        if route:
            all_done = all(o.status == 'Доставлен' for o in route.orders)
            if all_done and len(route.orders) > 0:
                route.status = 'Выполнен'
    db.session.commit()
    flash(f'✅ Заказ №{order.id} доставлен!', 'success')
    return redirect(url_for('driver_route', route_id=route_id) if route_id else url_for('driver_entry'))


@app.route('/driver/order/<int:order_id>/problem', methods=['POST'])
def driver_order_problem(order_id):
    order = Order.query.get_or_404(order_id)
    route_id = order.route_id
    comment = request.form.get('comment', '').strip()
    problem_text = f'⚠️ ПРОБЛЕМА: {comment}' if comment else '⚠️ ПРОБЛЕМА (без комментария)'
    if order.comment:
        order.comment = f'{order.comment} | {problem_text}'
    else:
        order.comment = problem_text
    db.session.commit()
    flash(f'⚠️ Проблема по заказу №{order.id} записана.', 'error')
    return redirect(url_for('driver_route', route_id=route_id) if route_id else url_for('driver_entry'))


@app.route('/driver/order/<int:order_id>/reset', methods=['POST'])
def driver_order_reset(order_id):
    order = Order.query.get_or_404(order_id)
    route_id = order.route_id
    order.status = 'В работе'
    if route_id:
        route = Route.query.get(route_id)
        if route and route.status == 'Выполнен':
            route.status = 'В работе'
    db.session.commit()
    flash(f'🔄 Заказ №{order.id} возвращён в работу.', 'success')
    return redirect(url_for('driver_route', route_id=route_id) if route_id else url_for('driver_entry'))


@app.route('/driver/route/<int:route_id>/finish', methods=['POST'])
def driver_route_finish(route_id):
    route = Route.query.get_or_404(route_id)
    route.status = 'Выполнен'
    db.session.commit()
    flash(f'🎉 Маршрут №{route.id} завершён!', 'success')
    return redirect(url_for('driver_home', driver_id=route.driver_id))


# ============================================================
# КЕГИ
# ============================================================
@app.route('/kegs')
def kegs():
    operation_filter = request.args.get('type', '').strip()
    client_filter = request.args.get('client_id', '').strip()
    date_from = request.args.get('date_from', '').strip()
    date_to = request.args.get('date_to', '').strip()

    query = Keg.query
    if operation_filter:
        query = query.filter(Keg.operation_type == operation_filter)
    if client_filter:
        try:
            query = query.filter(Keg.client_id == int(client_filter))
        except ValueError:
            pass
    if date_from:
        try:
            query = query.filter(Keg.operation_date >= datetime.strptime(date_from, '%Y-%m-%d').date())
        except ValueError:
            pass
    if date_to:
        try:
            query = query.filter(Keg.operation_date <= datetime.strptime(date_to, '%Y-%m-%d').date())
        except ValueError:
            pass

    kegs_list = query.order_by(Keg.operation_date.desc(), Keg.id.desc()).all()

    total_issued = db.session.query(func.sum(Keg.quantity)).filter(Keg.operation_type == 'Выдано').scalar() or 0
    total_returned = db.session.query(func.sum(Keg.quantity)).filter(Keg.operation_type == 'Возвращено').scalar() or 0
    stats = {
        'Выдано': total_issued,
        'Возвращено': total_returned,
        'Долг': total_issued - total_returned,
        'Сумма долга (₽)': (total_issued - total_returned) * KEG_PRICE,
    }

    all_clients = Client.query.order_by(Client.name).all()

    return render_template('kegs.html', kegs=kegs_list, stats=stats, clients=all_clients,
                           current_type=operation_filter, current_client=client_filter,
                           current_date_from=date_from, current_date_to=date_to)


@app.route('/kegs/add', methods=['GET', 'POST'])
def keg_add():
    if request.method == 'POST':
        client_id = request.form.get('client_id', '').strip()
        if not client_id:
            flash('Выберите клиента!', 'error')
            return redirect(url_for('keg_add'))
        operation_date = request.form.get('operation_date', '').strip()
        try:
            operation_date_val = datetime.strptime(operation_date, '%Y-%m-%d').date() if operation_date else date.today()
        except ValueError:
            operation_date_val = date.today()
        try:
            quantity = int(request.form.get('quantity', 0))
        except ValueError:
            quantity = 0
        if quantity <= 0:
            flash('Количество должно быть больше 0!', 'error')
            return redirect(url_for('keg_add'))

        keg = Keg(
            client_id=int(client_id),
            operation_date=operation_date_val,
            operation_type=request.form.get('operation_type', 'Выдано').strip(),
            quantity=quantity,
            comment=request.form.get('comment', '').strip(),
        )
        db.session.add(keg)
        db.session.commit()
        flash(f'Операция добавлена: {keg.operation_type} {quantity} кег', 'success')
        return redirect(url_for('kegs'))

    return render_template('keg_add.html',
                           clients=Client.query.order_by(Client.name).all(),
                           today=date.today().strftime('%Y-%m-%d'))


@app.route('/kegs/<int:keg_id>/delete', methods=['POST'])
def keg_delete(keg_id):
    keg = Keg.query.get_or_404(keg_id)
    db.session.delete(keg)
    db.session.commit()
    flash('Операция удалена!', 'success')
    return redirect(url_for('kegs'))


@app.route('/kegs/clients')
def kegs_clients():
    clients_list = Client.query.join(Keg, Keg.client_id == Client.id).distinct().all()
    result = []
    for client in clients_list:
        issued = db.session.query(func.sum(Keg.quantity)).filter(
            Keg.client_id == client.id, Keg.operation_type == 'Выдано'
        ).scalar() or 0
        returned = db.session.query(func.sum(Keg.quantity)).filter(
            Keg.client_id == client.id, Keg.operation_type == 'Возвращено'
        ).scalar() or 0
        debt = issued - returned
        if debt > 0 or issued > 0:
            result.append({
                'client': client,
                'issued': issued,
                'returned': returned,
                'debt': debt,
                'debt_sum': debt * KEG_PRICE,
            })
    result.sort(key=lambda x: x['debt'], reverse=True)
    total_debt = sum(r['debt'] for r in result)
    total_debt_sum = sum(r['debt_sum'] for r in result)
    return render_template('kegs_clients.html',
                           result=result, total_debt=total_debt,
                           total_debt_sum=total_debt_sum, keg_price=KEG_PRICE)


@app.route('/export/kegs')
def export_kegs():
    operation_filter = request.args.get('type', '').strip()
    client_filter = request.args.get('client_id', '').strip()
    date_from = request.args.get('date_from', '').strip()
    date_to = request.args.get('date_to', '').strip()

    query = Keg.query
    if operation_filter:
        query = query.filter(Keg.operation_type == operation_filter)
    if client_filter:
        try:
            query = query.filter(Keg.client_id == int(client_filter))
        except ValueError:
            pass
    if date_from:
        try:
            query = query.filter(Keg.operation_date >= datetime.strptime(date_from, '%Y-%m-%d').date())
        except ValueError:
            pass
    if date_to:
        try:
            query = query.filter(Keg.operation_date <= datetime.strptime(date_to, '%Y-%m-%d').date())
        except ValueError:
            pass

    kegs_list = query.order_by(Keg.operation_date.desc(), Keg.id.desc()).all()

    wb = Workbook()
    ws = wb.active
    ws.title = 'Операции с кегами'
    columns = ['№', 'Дата', 'Клиент', 'Операция', 'Количество', 'Комментарий']
    style_header(ws, 1, columns)
    for row_idx, keg in enumerate(kegs_list, 2):
        ws.cell(row=row_idx, column=1, value=keg.id)
        ws.cell(row=row_idx, column=2, value=keg.operation_date.strftime('%d.%m.%Y') if keg.operation_date else '')
        ws.cell(row=row_idx, column=3, value=keg.client.name if keg.client else '')
        ws.cell(row=row_idx, column=4, value=keg.operation_type)
        ws.cell(row=row_idx, column=5, value=keg.quantity)
        ws.cell(row=row_idx, column=6, value=keg.comment or '')
        for col_idx in range(1, 7):
            style_cell(ws.cell(row=row_idx, column=col_idx), 'center' if col_idx in [1, 2, 4, 5] else 'left')
    auto_width(ws)
    ws.freeze_panes = 'A2'
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    filename = f'Кеги_{datetime.now().strftime("%d.%m.%Y_%H-%M")}.xlsx'
    return send_file(output, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     as_attachment=True, download_name=filename)


@app.route('/export/kegs_clients')
def export_kegs_clients():
    clients_list = Client.query.join(Keg, Keg.client_id == Client.id).distinct().all()
    wb = Workbook()
    ws = wb.active
    ws.title = 'Долги по кегам'
    columns = ['#', 'Клиент', 'Телефон', 'Выдано', 'Возвращено', 'Долг', 'Сумма долга (₽)']
    style_header(ws, 1, columns)
    row_idx = 2
    total_debt = 0
    total_sum = 0
    for client in clients_list:
        issued = db.session.query(func.sum(Keg.quantity)).filter(
            Keg.client_id == client.id, Keg.operation_type == 'Выдано'
        ).scalar() or 0
        returned = db.session.query(func.sum(Keg.quantity)).filter(
            Keg.client_id == client.id, Keg.operation_type == 'Возвращено'
        ).scalar() or 0
        debt = issued - returned
        if debt > 0:
            ws.cell(row=row_idx, column=1, value=row_idx - 1)
            ws.cell(row=row_idx, column=2, value=client.name)
            ws.cell(row=row_idx, column=3, value=client.phone or '')
            ws.cell(row=row_idx, column=4, value=issued)
            ws.cell(row=row_idx, column=5, value=returned)
            ws.cell(row=row_idx, column=6, value=debt)
            ws.cell(row=row_idx, column=7, value=debt * KEG_PRICE)
            for col_idx in range(1, 8):
                style_cell(ws.cell(row=row_idx, column=col_idx), 'center' if col_idx != 2 else 'left')
            total_debt += debt
            total_sum += debt * KEG_PRICE
            row_idx += 1
    if total_debt > 0:
        ws.cell(row=row_idx, column=1, value='ИТОГО:').font = Font(bold=True)
        ws.cell(row=row_idx, column=6, value=total_debt).font = Font(bold=True)
        ws.cell(row=row_idx, column=7, value=total_sum).font = Font(bold=True)
    auto_width(ws)
    ws.freeze_panes = 'A2'
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    filename = f'Долги_по_кегам_{datetime.now().strftime("%d.%m.%Y_%H-%M")}.xlsx'
    return send_file(output, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     as_attachment=True, download_name=filename)


# ============================================================
# ИМПОРТ
# ============================================================
@app.route('/import', methods=['GET', 'POST'])
def import_excel():
    if request.method == 'POST':
        file = request.files.get('file')
        order_date_str = request.form.get('order_date', '').strip()
        ignore_duplicates = request.form.get('ignore_duplicates') == 'on'

        if not file or file.filename == '':
            flash('Выберите файл для загрузки!', 'error')
            return redirect(url_for('import_excel'))
        if not file.filename.endswith(('.xls', '.xlsx')):
            flash('Файл должен быть в формате .xls или .xlsx', 'error')
            return redirect(url_for('import_excel'))

        try:
            order_date_val = datetime.strptime(order_date_str, '%Y-%m-%d').date() if order_date_str else date.today()
        except ValueError:
            order_date_val = date.today()

        filename = secure_filename(file.filename)
        filepath = os.path.join(app.config['UPLOAD_FOLDER'], filename)
        file.save(filepath)
        file_hash = get_file_hash(filepath)

        existing_import = ImportLog.query.filter_by(file_hash=file_hash, order_date=order_date_val).first()
        if existing_import and not ignore_duplicates:
            flash('⚠️ Этот файл уже загружался. Импорт пропущен.', 'error')
            return redirect(url_for('import_excel'))

        try:
            df = pd.read_excel(filepath, header=0)
        except Exception as e:
            flash(f'Не удалось прочитать файл: {e}', 'error')
            return redirect(url_for('import_excel'))

        if len(df.columns) < 3:
            flash('В файле должно быть минимум 3 колонки', 'error')
            return redirect(url_for('import_excel'))

        df.columns = ['Клиент', 'Адрес', 'Сумма'] + list(df.columns[3:])

        clients_created = 0
        clients_found = 0
        orders_created = 0
        orders_skipped = 0
        total_amount = 0

        for idx, row in df.iterrows():
            name = str(row['Клиент']).strip() if pd.notna(row['Клиент']) else ''
            address = str(row['Адрес']).strip() if pd.notna(row['Адрес']) else ''
            if not name or name.lower() in ['nan', 'контрагенты', 'итого', 'клиент']:
                continue
            try:
                amount = float(row['Сумма']) if pd.notna(row['Сумма']) else 0
            except (ValueError, TypeError):
                amount = 0
            if amount <= 0:
                continue

            client = find_client_by_name(name)
            if client:
                clients_found += 1
                if not client.address and address:
                    client.address = address
            else:
                client = Client(name=name, address=address, group_abc='C',
                                comment=f'Импортирован из файла {filename}')
                db.session.add(client)
                db.session.flush()
                clients_created += 1

            if not ignore_duplicates:
                if order_exists(client.id, amount, order_date_val):
                    orders_skipped += 1
                    continue

            order = Order(client_id=client.id, order_date=order_date_val, amount=amount,
                          status='Новый',
                          comment=f'Импорт из файла: {filename}' + (f' | Адрес: {address}' if address else ''))
            db.session.add(order)
            orders_created += 1
            total_amount += amount

        log = ImportLog(filename=filename, file_hash=file_hash, order_date=order_date_val,
                        clients_created=clients_created, clients_found=clients_found,
                        orders_created=orders_created, orders_skipped=orders_skipped,
                        total_amount=total_amount)
        db.session.add(log)

        try:
            db.session.commit()
            msg = f'✅ Импорт завершён! Создано клиентов: {clients_created}, найдено: {clients_found}, заказов: {orders_created}'
            if orders_skipped > 0:
                msg += f', пропущено дублей: {orders_skipped}'
            msg += f', на сумму: {total_amount:,.0f} ₽'
            flash(msg, 'success')
        except Exception as e:
            db.session.rollback()
            flash(f'Ошибка: {e}', 'error')
        return redirect(url_for('import_excel'))

    history = ImportLog.query.order_by(ImportLog.created_at.desc()).limit(20).all()
    return render_template('import.html', today=date.today().strftime('%Y-%m-%d'), history=history)


@app.route('/import/clear_history', methods=['POST'])
def import_clear_history():
    ImportLog.query.delete()
    db.session.commit()
    flash('История импортов очищена.', 'success')
    return redirect(url_for('import_excel'))

# ============================================================
# ДАШБОРД
# ============================================================
@app.route('/dashboard')
def dashboard():
    stats = {
        'clients': Client.query.count(),
        'orders': Order.query.count(),
        'drivers': Driver.query.count(),
        'routes': Route.query.count(),
    }

    orders_by_status = {
        'Новый': Order.query.filter_by(status='Новый').count(),
        'В работе': Order.query.filter_by(status='В работе').count(),
        'Доставлен': Order.query.filter_by(status='Доставлен').count(),
        'Отменён': Order.query.filter_by(status='Отменён').count(),
    }

    sums_by_status = {
        'Новый': Order.query.filter_by(status='Новый').with_entities(func.sum(Order.amount)).scalar() or 0,
        'В работе': Order.query.filter_by(status='В работе').with_entities(func.sum(Order.amount)).scalar() or 0,
        'Доставлен': Order.query.filter_by(status='Доставлен').with_entities(func.sum(Order.amount)).scalar() or 0,
        'Отменён': Order.query.filter_by(status='Отменён').with_entities(func.sum(Order.amount)).scalar() or 0,
    }

    total_amount = sum(sums_by_status.values())

    top_clients = db.session.query(
        Client.name, func.sum(Order.amount).label('total')
    ).join(Order).group_by(Client.id).order_by(func.sum(Order.amount).desc()).limit(5).all()

    active_routes = Route.query.filter(Route.status != 'Выполнен').order_by(Route.route_date.desc()).all()
    recent_orders = Order.query.order_by(Order.created_at.desc()).limit(10).all()
    free_orders_count = Order.query.filter(Order.route_id.is_(None)).count()

    today = date.today()
    dates_30 = [(today - timedelta(days=i)) for i in range(29, -1, -1)]
    chart_dates = [d.strftime('%d.%m') for d in dates_30]
    chart_sales = []
    for d in dates_30:
        total = db.session.query(func.sum(Order.amount)).filter(Order.order_date == d).scalar() or 0
        chart_sales.append(round(total, 0))

    chart_status_labels = ['Новый', 'В работе', 'Доставлен', 'Отменён']
    chart_status_values = [
        orders_by_status['Новый'], orders_by_status['В работе'],
        orders_by_status['Доставлен'], orders_by_status['Отменён'],
    ]

    chart_clients_labels = [c.name[:25] for c in top_clients]
    chart_clients_values = [round(c.total, 0) for c in top_clients]

    drivers_chart = db.session.query(
        Driver.name, func.sum(Order.amount).label('total')
    ).outerjoin(Route, Route.driver_id == Driver.id)\
     .outerjoin(Order, Order.route_id == Route.id)\
     .group_by(Driver.id).order_by(func.sum(Order.amount).desc()).limit(5).all()

    chart_drivers_labels = [d.name[:20] for d in drivers_chart]
    chart_drivers_values = [round(d.total or 0, 0) for d in drivers_chart]

    keg_debt_total = (db.session.query(func.sum(Keg.quantity)).filter(Keg.operation_type == 'Выдано').scalar() or 0) - \
                     (db.session.query(func.sum(Keg.quantity)).filter(Keg.operation_type == 'Возвращено').scalar() or 0)

    clients_no_coords = Client.query.filter(
        (Client.lat.is_(None)) | (Client.lon.is_(None))
    ).count()

    return render_template('dashboard.html',
                           stats=stats, orders_by_status=orders_by_status,
                           sums_by_status=sums_by_status, total_amount=total_amount,
                           top_clients=top_clients, active_routes=active_routes,
                           recent_orders=recent_orders, free_orders_count=free_orders_count,
                           chart_dates=chart_dates, chart_sales=chart_sales,
                           chart_status_labels=chart_status_labels, chart_status_values=chart_status_values,
                           chart_clients_labels=chart_clients_labels, chart_clients_values=chart_clients_values,
                           chart_drivers_labels=chart_drivers_labels, chart_drivers_values=chart_drivers_values,
                           keg_debt_total=keg_debt_total,
                           clients_no_coords=clients_no_coords)


@app.route('/export/dashboard')
def export_dashboard():
    wb = Workbook()
    ws1 = wb.active
    ws1.title = 'Сводка'
    ws1.merge_cells('A1:C1')
    title = ws1.cell(row=1, column=1, value='СВОДНЫЙ ОТЧЁТ')
    title.font = Font(bold=True, size=16)
    title.alignment = Alignment(horizontal='center', vertical='center')
    ws1.row_dimensions[1].height = 25
    ws1.cell(row=2, column=1, value=f'Дата: {datetime.now().strftime("%d.%m.%Y %H:%M")}').font = Font(italic=True, color='666666')
    row = 4
    ws1.cell(row=row, column=1, value='ПОКАЗАТЕЛЬ').font = Font(bold=True, color='FFFFFF')
    ws1.cell(row=row, column=2, value='ЗНАЧЕНИЕ').font = Font(bold=True, color='FFFFFF')
    for col in [1, 2]:
        ws1.cell(row=row, column=col).fill = PatternFill(start_color='667EEA', end_color='667EEA', fill_type='solid')
    data = [
        ('Клиентов', Client.query.count()),
        ('Водителей', Driver.query.count()),
        ('Маршрутов', Route.query.count()),
        ('Заказов всего', Order.query.count()),
        ('', ''),
        ('Общая сумма (₽)', Order.query.with_entities(func.sum(Order.amount)).scalar() or 0),
        ('Заказы без маршрута', Order.query.filter(Order.route_id.is_(None)).count()),
        ('Клиентов без координат', Client.query.filter((Client.lat.is_(None)) | (Client.lon.is_(None))).count()),
    ]
    for row_idx, (label, value) in enumerate(data, 5):
        ws1.cell(row=row_idx, column=1, value=label)
        ws1.cell(row=row_idx, column=2, value=value)
        style_cell(ws1.cell(row=row_idx, column=1), 'left')
        style_cell(ws1.cell(row=row_idx, column=2), 'center')
    ws1.column_dimensions['A'].width = 30
    ws1.column_dimensions['B'].width = 20

    ws2 = wb.create_sheet('Топ-клиенты')
    columns = ['#', 'Клиент', 'Сумма заказов (₽)', 'Количество заказов']
    style_header(ws2, 1, columns)
    top_clients_data = db.session.query(
        Client.name, func.sum(Order.amount).label('total'), func.count(Order.id).label('cnt')
    ).join(Order).group_by(Client.id).order_by(func.sum(Order.amount).desc()).limit(20).all()
    for row_idx, (name, total, cnt) in enumerate(top_clients_data, 2):
        ws2.cell(row=row_idx, column=1, value=row_idx - 1)
        ws2.cell(row=row_idx, column=2, value=name)
        ws2.cell(row=row_idx, column=3, value=total)
        ws2.cell(row=row_idx, column=4, value=cnt)
        for col_idx in range(1, 5):
            style_cell(ws2.cell(row=row_idx, column=col_idx), 'center' if col_idx in [1, 3, 4] else 'left')
    auto_width(ws2)

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    filename = f'Отчёт_{datetime.now().strftime("%d.%m.%Y_%H-%M")}.xlsx'
    return send_file(output, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     as_attachment=True, download_name=filename)

# ============================================================
# ЗАПУСК
# ============================================================
# ============================================================
# ЗАПУСК (продакшн-режим через waitress)
# ============================================================
def _print_startup_info(host, port):
    import socket
    print("=" * 60)
    print("🚀 ЗАПУСК ПРИЛОЖЕНИЯ")
    print("=" * 60)

    # Проверка ключа DaData
    try:
        import geocode as _gc
        if _gc.DADATA_API_KEY:
            print(f"✅ DADATA_API_KEY: {_gc.DADATA_API_KEY[:8]}...")
        else:
            print("❌ DADATA_API_KEY НЕ НАЙДЕН (проверь secret.py)")
    except Exception as e:
        print(f"⚠️ Не удалось проверить ключ: {e}")

    print("-" * 60)
    print(f"📍 Локальный доступ:   http://127.0.0.1:{port}")
    print(f"📱 Мобильная версия:   http://127.0.0.1:{port}/driver")

    # Все локальные IP этого компьютера
    try:
        hostname = socket.gethostname()
        local_ips = set()
        for info in socket.getaddrinfo(hostname, None):
            ip = info[4][0]
            if ':' not in ip and not ip.startswith('127.'):
                local_ips.add(ip)
        if local_ips:
            print("-" * 60)
            print("🌐 Доступ по локальной сети (для других устройств в Wi-Fi):")
            for ip in sorted(local_ips):
                print(f"   http://{ip}:{port}")
                print(f"   http://{ip}:{port}/driver   ← для водителей")
    except Exception:
        pass

    print("=" * 60)
    print("⚠️  Чтобы остановить — нажми Ctrl+C")
    print("=" * 60)


if __name__ == '__main__':
    # host='0.0.0.0' — слушать все сетевые интерфейсы (доступ из LAN)
    # port=5000 — можно сменить, если занят
    HOST = '0.0.0.0'
    PORT = 5000

    _print_startup_info(HOST, PORT)

    # Пытаемся использовать waitress (продакшн-сервер).
    # Если он не установлен — падаем обратно на встроенный Flask с предупреждением.
    try:
        from waitress import serve
        print("✅ Используется waitress (продакшн-режим)")
        serve(app, host=HOST, port=PORT, threads=8)
    except ImportError:
        print("⚠️ waitress не установлен — работаю на встроенном сервере Flask.")
        print("   Установить:  pip install waitress")
        print("   Встроенный сервер НЕ предназначен для продакшена, работает медленно.")
        app.run(host=HOST, port=PORT, debug=False)